import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

os.makedirs("docs/word_documents", exist_ok=True)

def create_base_doc(title, subtitle):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    r_sub = p_title.add_run(subtitle.upper() + "\n")
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(37, 99, 235)

    r_title = p_title.add_run(title)
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    return doc

def add_meta_table(doc, marks="4 Marks"):
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = 'Table Grid'
    data = [
        ("Date", "20 October 2026"),
        ("Team ID", "HP-TEAM-2026"),
        ("Project Name", "HotelPro — Hotel Booking and Management Platform with AI Assistant"),
        ("Maximum Marks", marks)
    ]
    for i, (k, v) in enumerate(data):
        row = meta_table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        shd0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        row.cells[0]._tc.get_or_add_tcPr().append(shd0)
        for p in row.cells[0].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9.5)
        for p in row.cells[1].paragraphs:
            for run in p.runs:
                run.font.size = Pt(9.5)
    doc.add_paragraph()

def add_table(doc, headers, data):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        hdr.cells[i].text = h
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
        hdr.cells[i]._tc.get_or_add_tcPr().append(shd)
        for p in hdr.cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for r_idx, row_values in enumerate(data):
        row = table.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row.cells[c_idx].text = str(val)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            row.cells[c_idx]._tc.get_or_add_tcPr().append(shd)
            for p in row.cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    doc.add_paragraph()

# -------------------------------------------------------------
# 1. Empathy Map Canvas
# -------------------------------------------------------------
def build_01():
    doc = create_base_doc("Empathize & Discover — Empathy Map Canvas", "Phase 1: Ideation Phase")
    add_meta_table(doc, "4 Marks")
    
    doc.add_heading("1. Purpose of the Empathy Map Canvas", level=2)
    doc.add_paragraph(
        "To deeply understand the real-life behaviors, daily frustrations, and emotional triggers of both travellers "
        "and independent hotel managers, guiding the design of an overbooking-proof, transparent hotel booking platform."
    )
    
    doc.add_heading("2. Empathy Map: The Traveller (Aarav, 24 / Meera, 38)", level=2)
    guest_data = [
        ["1. Who are we empathizing with?", "Aarav (24, young professional planning a trip with friends) & Meera (38, mother booking for family of 5)."],
        ["2. What do they need to DO?", "Book verified rooms without fear of arriving to a double-booked hotel; see full price upfront; cancel easily if plans change."],
        ["3. What do they SEE?", "Aggregator sites with misleading base prices, pop-ups claiming 'only 1 room left', and outdated photos taken a decade ago."],
        ["4. What do they SAY?", "'Is this the actual final price or are taxes added at the end?' 'I hope they don't sell my room to a walk-in guest before I arrive.'"],
        ["5. What do they DO?", "Keep 5 booking tabs open, screenshot options on WhatsApp, rush anxiously through payment forms, and read Reddit for dirty room warnings."],
        ["6. What do they HEAR?", "Friends warning about midnight check-in rejections, customer care stating 'this was an aggregator glitch', and rigid no-refund policies."],
        ["7. PAINS (Frustrations)", "Overbooking panic at late night; hidden taxes (GST/convenience fees) suddenly added on the payment screen; lack of clear pet/pool rules."],
        ["8. GAINS (Desires)", "Guaranteed reservation backed by an atomic 10-minute hold window; upfront 12% GST breakdown; instant AI answers to hotel policy questions."]
    ]
    add_table(doc, ["Dimension", "Observations & Insights"], guest_data)

    doc.add_heading("3. Empathy Map: The Independent Hotelier (Rakesh, 45)", level=2)
    staff_data = [
        ["1. Who are we empathizing with?", "Rakesh (45), runs a 20-room boutique hotel in Calangute, North Goa."],
        ["2. What do they need to DO?", "Maintain high occupancy, avoid double-booking disputes, manage daily guest arrivals, and reply directly to customer reviews."],
        ["3. What do they SEE?", "Front desk staff juggling paper diaries, phone bookings, and OTA portals; families arguing in the lobby over lost reservations."],
        ["4. What do they SAY?", "'We don't need complicated 5-star enterprise software. We just need to know who is arriving today and which rooms are empty.'"],
        ["5. What do they DO?", "Check rooms personally, write reservations by hand in notebooks, and struggle with 20% to 25% commission fees from dominant aggregators."],
        ["6. What do they HEAR?", "Front desk staff saying: 'Sir, another booking just arrived online while I was taking a phone reservation.'"],
        ["7. PAINS", "Reputation destruction from technical overbookings; exorbitant commission charges; clunky software that staff refuse to use."],
        ["8. GAINS", "A 1-click Check-in/Check-out web portal; atomic concurrency lock guaranteeing zero double-bookings; direct inventory management."]
    ]
    add_table(doc, ["Dimension", "Observations & Insights"], staff_data)
    
    doc.save("docs/word_documents/01_Empathy_Map_Canvas.docx")
    print("Generated: 01_Empathy_Map_Canvas.docx")

# -------------------------------------------------------------
# 2. Define Problem Statements
# -------------------------------------------------------------
def build_02():
    doc = create_base_doc("Customer Problem Statement Template", "Phase 1: Ideation Phase")
    add_meta_table(doc, "2 Marks")

    doc.add_heading("Problem Statement Formulation", level=2)
    doc.add_paragraph("Using the Miro Customer Problem Statement framework (I am / I'm trying to / But / Because / Which makes me feel):")

    ps_data = [
        [
            "PS-1",
            "A working professional planning a family vacation to Goa.",
            "Secure a guaranteed room reservation online for specific holiday dates.",
            "When I arrive at the hotel lobby after an exhausting journey, the front desk tells me the room was given away or booked twice.",
            "The reservation platform allows concurrent checkouts without database-level concurrency locks or atomic room holds.",
            "Stranded, humiliated, and deeply anxious about finding emergency shelter for my family at midnight."
        ],
        [
            "PS-2",
            "A budget-conscious young traveller.",
            "Compare hotel prices transparently and confirm pet and amenity policies.",
            "The final checkout amount jumps by 30% to 40% due to surprise taxes, and property details like pool access or pet rules remain ambiguous.",
            "Traditional platforms obscure mandatory fees until the final card swipe and rely on static, unverified text blurbs instead of real-time catalog intelligence.",
            "Cheated, distrustful, and frustrated by having to open 10 tabs and make repeated phone calls to the front desk."
        ],
        [
            "PS-3",
            "An independent hotel owner managing a 20-room boutique resort.",
            "Keep rooms occupied, manage daily arrivals, and update prices easily.",
            "Rooms frequently get double-booked during holiday rushes, and front-desk staff struggle with clumsy paper registers or bloated enterprise software.",
            "We lack a lightweight, cloud-based management portal with live room status, atomic hold synchronization, and 1-click check-in/out workflows.",
            "Overwhelmed, vulnerable to bad online reviews, and constantly worried about operational chaos damaging our livelihood."
        ],
        [
            "PS-4",
            "A busy business traveller with tight schedules.",
            "Quickly find a hotel room matching specific natural constraints (e.g., 'quiet room near Whitefield Bengaluru under 4000 with high-speed WiFi').",
            "I have to spend 25 minutes fiddling with dozens of nested filter checkboxes and scrolling through irrelevant properties.",
            "Search engines rely on rigid keyword lookups and lack conversational AI grounded in live room availability and verified location data.",
            "Impatient, fatigued, and dissatisfied with the cumbersome booking experience."
        ]
    ]
    add_table(doc, ["PS No.", "I am (Customer)", "I'm trying to", "But", "Because", "Which makes me feel"], ps_data)

    doc.save("docs/word_documents/02_Define_Problem_Statements.docx")
    print("Generated: 02_Define_Problem_Statements.docx")

# -------------------------------------------------------------
# 3. Brainstorm & Idea Prioritization
# -------------------------------------------------------------
def build_03():
    doc = create_base_doc("Brainstorm & Idea Prioritization Template", "Phase 1: Ideation Phase")
    add_meta_table(doc, "4 Marks")

    doc.add_heading("Step 1: Selection of Problem Statement", level=2)
    doc.add_paragraph(
        "Selected Core Challenge: How might we build a hotel reservation platform that guarantees zero overbooking under high concurrent traffic, "
        "provides transparent upfront pricing, and integrates a grounded AI concierge that answers traveller questions using live catalog records?"
    )

    doc.add_heading("Step 2: Idea Listing & Grouping", level=2)
    ideas = [
        ["Group A: Overbooking Prevention", "Atomic 10-minute hold window; real-time date overlap query; automated background hold-expiry worker; max 3 unpaid holds per user."],
        ["Group B: Pricing & Payments", "Authoritative server-side price computation; transparent 12% GST itemization; Razorpay test sandbox with HMAC verification; tiered 48h/24h refund policy."],
        ["Group C: Grounded AI Concierge", "RAG assistant grounded strictly in HotelPro catalog; natural language search parser; AI review synthesizer for pros/cons."],
        ["Group D: Hotelier Operations", "Lightweight front desk portal; 1-click Check-in/Check-out; direct public manager replies to reviews; room blackout dates."],
        ["Group E: Security & Lifecycle", "Dual-token authentication (15m access + 7d httpOnly refresh cookie); 5-attempt brute force lockout; verified post-stay reviews only."]
    ]
    add_table(doc, ["Functional Group", "Brainstormed Solutions & Architectural Mechanisms"], ideas)

    doc.add_heading("Step 3: MoSCoW Prioritization", level=2)
    moscow = [
        ["Must Have (MVP)", "Atomic 10-minute hold window; date overlap algorithm; dual-token auth; Razorpay checkout; server-calculated 12% GST; My Bookings dashboard."],
        ["Should Have", "Grounded RAG AI concierge; front desk Check-in/Check-out portal; admin revenue analytics; verified reviews; coupon code engine."],
        ["Could Have (Stand-out)", "Natural language search parsing; AI review summary (pros/cons); public manager replies; active session revocation."],
        ["Won't Have (v1)", "Native iOS/Android app; multi-currency conversions; real-time WebSocket chat."]
    ]
    add_table(doc, ["MoSCoW Category", "Features & Engineering Deliverables"], moscow)

    doc.save("docs/word_documents/03_Brainstorm_and_Idea_Prioritization.docx")
    print("Generated: 03_Brainstorm_and_Idea_Prioritization.docx")

# -------------------------------------------------------------
# 4. Project Planning
# -------------------------------------------------------------
def build_04():
    doc = create_base_doc("Project Planning Template (Product Backlog & Sprints)", "Phase 2: Project Planning Phase")
    add_meta_table(doc, "5 Marks")

    backlog = [
        ["Sprint-1", "User Authentication", "USN-01", "Register with email, password strength check & bcrypt hash", "3", "High", "Backend"],
        ["Sprint-1", "User Authentication", "USN-02", "Dual-token login (15m access + 7d httpOnly refresh cookie)", "3", "High", "Full Stack"],
        ["Sprint-1", "Security", "USN-03", "Brute-force protection: 15-min lockout after 5 wrong attempts", "2", "High", "Backend"],
        ["Sprint-2", "Hotel Discovery", "USN-05", "Search by destination city, dates, and live room availability", "3", "High", "Full Stack"],
        ["Sprint-2", "Hotel Discovery", "USN-06", "Filter by price slider, star rating, property type & amenities", "3", "High", "Frontend"],
        ["Sprint-2", "Hotel Details", "USN-07", "Photo gallery, room categories, policies & verified reviews", "2", "High", "Frontend"],
        ["Sprint-3", "Overbooking Engine", "USN-08", "Atomic 10-minute room hold with live countdown timer", "5", "Critical", "Backend"],
        ["Sprint-3", "Pricing & Billing", "USN-09", "Server-calculated breakdown: base price + 12% GST - coupon", "3", "High", "Full Stack"],
        ["Sprint-3", "Payments", "USN-10", "Razorpay checkout simulation with HMAC-SHA256 verification", "5", "Critical", "Full Stack"],
        ["Sprint-3", "Booking Lifecycle", "USN-11", "My Bookings hub with automated 48h/24h cancellation refunds", "3", "High", "Full Stack"],
        ["Sprint-4", "Staff Operations", "USN-12", "Hotelier portal: daily arrivals and 1-click Check-in/Check-out", "3", "High", "Full Stack"],
        ["Sprint-4", "Admin Governance", "USN-14", "Admin dashboard: revenue KPIs, occupancy % & user role control", "3", "High", "Full Stack"],
        ["Sprint-4", "AI Concierge", "USN-15", "Catalog-grounded RAG assistant with Google Gemini integration", "5", "High", "AI / Backend"]
    ]
    add_table(doc, ["Sprint", "Epic", "Story ID", "User Story / Task", "Points", "Priority", "Role"], backlog)

    doc.save("docs/word_documents/04_Project_Planning_and_Sprint_Schedule.docx")
    print("Generated: 04_Project_Planning_and_Sprint_Schedule.docx")

# -------------------------------------------------------------
# 5. Solution Requirements
# -------------------------------------------------------------
def build_05():
    doc = create_base_doc("Solution Requirements (Functional & Non-Functional)", "Phase 3: Requirement Analysis Phase")
    add_meta_table(doc, "4 Marks")

    fr = [
        ["FR-1", "User Authentication", "Registration with password strength validation; dual-token login (15m JWT + 7d httpOnly cookie); 5-attempt brute-force lock."],
        ["FR-2", "Hotel Search & Discovery", "Multi-attribute search (city, dates, guests); date overlap algorithm excluding sold-out rooms; price/star/amenity filters."],
        ["FR-3", "Property Details & Gallery", "High-res gallery; room categories with bed types and capacity; verified check-in/out and pet policies."],
        ["FR-4", "Overbooking-Proof Engine", "Atomic 10-minute hold with live countdown timer; background cron worker auto-expiring unpaid holds every 60s."],
        ["FR-5", "Payments & Server Pricing", "Server-decided price breakdown with 12% GST; Razorpay test checkout with HMAC verification; printable voucher."],
        ["FR-6", "Booking Lifecycle", "My Bookings tabs (Upcoming, Past, Cancelled); automated refund engine (>48h: 100%, 24-48h: 50%, <24h: 0%)."],
        ["FR-7", "Staff Operations", "Hotelier front desk portal; 1-click Check-in and Check-out status updater; manager replies to verified reviews."],
        ["FR-8", "Admin Governance", "Executive analytics (revenue, bookings, occupancy rate); user role management (`guest`, `staff`, `admin`)."],
        ["FR-9", "Grounded AI Concierge", "Conversational assistant grounded in catalog records; natural language search parsing; review pros/cons synthesis."]
    ]
    add_table(doc, ["FR No.", "Functional Epic", "Detailed Functional Requirement"], fr)

    nfr = [
        ["NFR-1", "Usability & Mobile-First", "Clean, modern responsive interface; WCAG AA compliant contrast; intuitive hold countdown timer."],
        ["NFR-2", "Zero-Overbooking Guarantee", "Atomic database operations prevent double-booking even under 20 concurrent requests for the last room."],
        ["NFR-3", "Security & Data Protection", "Bcrypt password hashing; tokens in secure httpOnly cookies; Zod input validation; Helmet security headers."],
        ["NFR-4", "Performance & Latency", "Search pages load in under 3 seconds on 4G; API responses resolve in under 500ms; database indexed on city/dates."],
        ["NFR-5", "AI Safety & Guardrails", "AI assistant grounded strictly in HotelPro catalog; zero hallucination of room rates or availability; polite deflection of off-topic queries."]
    ]
    add_table(doc, ["NFR No.", "Quality Attribute", "Specification & Target Metric"], nfr)

    doc.save("docs/word_documents/05_Solution_Requirements_Functional_and_NonFunctional.docx")
    print("Generated: 05_Solution_Requirements_Functional_and_NonFunctional.docx")

# -------------------------------------------------------------
# 6. Technology Stack
# -------------------------------------------------------------
def build_06():
    doc = create_base_doc("Technology Stack (Architecture & Stack)", "Phase 3: Requirement Analysis Phase")
    add_meta_table(doc, "4 Marks")

    t1 = [
        ["1", "User Interface (Frontend)", "Single-page responsive web client for travellers and hoteliers", "React 18, Vite, Tailwind CSS, React Router v7"],
        ["2", "Icons & Components", "Lightweight vector icons matching modern travel portals", "Lucide React"],
        ["3", "HTTP & State", "API client with transparent 401 token refresh interceptors", "Axios"],
        ["4", "Backend Web API", "Modular REST API implementing business rules and hold engine", "Node.js, Express.js"],
        ["5", "Request Validation", "Type-safe runtime schema validation for incoming payloads", "Zod"],
        ["6", "Security & Session", "Dual-token session handling and password hashing", "bcryptjs, jsonwebtoken, cookie-parser, helmet"],
        ["7", "Database & Persistence", "Document storage for hotels, rooms, bookings, and sessions", "MongoDB Atlas / Mongoose"],
        ["8", "Payment Simulation", "Test checkout simulation for UPI, Cards, and Net Banking", "Razorpay Sandbox (Test Mode)"],
        ["9", "Artificial Intelligence", "Catalog-grounded conversational assistant & NLP parsing", "Google Gemini API (gemini-1.5-flash)"]
    ]
    add_table(doc, ["S.No", "Component", "Description", "Technology Selected"], t1)

    t2 = [
        ["1", "Open Source Stack", "100% open-source components with zero vendor lock-in", "React, Node.js, Express, Tailwind, Mongoose"],
        ["2", "Concurrency & Scalability", "Stateless API with database-level isolation", "JWT tokens + MongoDB compound indexes for date overlap"],
        ["3", "Automated Testing", "Native Node.js test runner verifying core business rules", "node --test tests/**/*.test.js (sub-second execution)"],
        ["4", "Cloud Deployment Tier", "Free-tier hosting compliance for academic evaluation", "Vercel (Frontend), Render (Backend), MongoDB Atlas (M0 Free)"]
    ]
    add_table(doc, ["S.No", "Characteristic", "Architectural Strategy", "Implementation Details"], t2)

    doc.save("docs/word_documents/06_Technology_Stack_and_Architecture.docx")
    print("Generated: 06_Technology_Stack_and_Architecture.docx")

# -------------------------------------------------------------
# 7. Data Flow Diagrams & User Stories
# -------------------------------------------------------------
def build_07():
    doc = create_base_doc("Data Flow Diagram & User Stories", "Phase 3: Requirement Analysis Phase")
    add_meta_table(doc, "4 Marks")

    doc.add_heading("1. DFD Level 0 (Context Level)", level=2)
    doc.add_paragraph(
        "Entities: Guest / Traveller, Hotel Staff / Manager, Platform Administrator.\n"
        "Core System: HotelPro Core Booking & AI Engine.\n"
        "External Services: Razorpay Test Gateway, Google Gemini AI Platform.\n"
        "Flows: Search queries, atomic room holds, test payment payloads, arrivals tracking, and AI catalog queries."
    )

    doc.add_heading("2. User Stories Table", level=2)
    stories = [
        ["Guest", "Authentication", "As a guest, I can register and log in securely.", "Password validated (Aa1); JWT issued; httpOnly cookie set.", "High", "Sprint-1"],
        ["Guest", "Hotel Search", "As a traveller, I can search hotels by city, dates, and guests.", "Only hotels with real vacant rooms returned; past dates blocked.", "High", "Sprint-2"],
        ["Guest", "Hotel Filters", "As a traveller, I can filter by price, star rating, and amenities.", "Instant UI filtering; search state saved in shareable URL.", "High", "Sprint-2"],
        ["Guest", "Hold Engine", "As a guest, I want a 10-minute hold window during checkout.", "Room availability decrements; countdown timer visible.", "Critical", "Sprint-3"],
        ["Guest", "Server Pricing", "As a guest, I want an itemized breakdown with 12% GST.", "Calculated on server; transparently shows subtotal and taxes.", "High", "Sprint-3"],
        ["Guest", "Payments", "As a guest, I can pay via Razorpay in test mode.", "Razorpay modal opens; HMAC signature verified; voucher generated.", "Critical", "Sprint-3"],
        ["Guest", "Cancellations", "As a guest, I can cancel an upcoming stay with refund.", "Refund policy enforced (>48h: 100%, 24-48h: 50%, <24h: 0%).", "High", "Sprint-3"],
        ["Staff", "Front Desk", "As a manager, I can mark guests as Checked-in and Completed.", "Status updates with 1-click in the staff portal.", "High", "Sprint-4"],
        ["Admin", "Governance", "As an admin, I can view gross revenue and occupancy metrics.", "Dashboard displays revenue charts and user role controls.", "High", "Sprint-4"],
        ["User", "AI Concierge", "As a user, I can ask the AI questions about hotel amenities.", "AI answers accurately from catalog without hallucinating.", "High", "Sprint-4"]
    ]
    add_table(doc, ["User Type", "Epic", "User Story / Task", "Acceptance Criteria", "Priority", "Release"], stories)

    doc.save("docs/word_documents/07_Data_Flow_Diagrams_and_User_Stories.docx")
    print("Generated: 07_Data_Flow_Diagrams_and_User_Stories.docx")

# -------------------------------------------------------------
# 8. Problem - Solution Fit
# -------------------------------------------------------------
def build_08():
    doc = create_base_doc("Problem – Solution Fit Template", "Phase 4: Project Design Phase")
    add_meta_table(doc, "2 Marks")

    psf = [
        ["1", "Problem Statement", "Travellers face late-night overbooking disasters and surprise taxes at checkout. Independent boutique hoteliers suffer from manual double-booking chaos and 20-25% aggregator fees."],
        ["2", "Idea / Solution Description", "HotelPro: A full-stack MERN platform featuring an overbooking-proof reservation engine, 10-minute price hold, server-calculated 12% GST, Razorpay test payments, and a grounded Gemini AI concierge."],
        ["3", "Novelty / Uniqueness", "Atomic database hold locks; RAG AI assistant grounded strictly in hotel catalog data; transparent server-side billing; automated 48h/24h cancellation refunds."],
        ["4", "Social Impact", "Restores traveller peace of mind; empowers boutique hoteliers with free enterprise-grade front desk tools; eliminates deceptive drip pricing."],
        ["5", "Business Model", "Transparent 5% to 8% platform fee (vs 20-25% OTA fees); optional featured hotel placement; multi-property boutique chain licensing."],
        ["6", "Scalability", "Stateless microservice-ready Node.js API; horizontal MongoDB Atlas sharding; pre-architected for WebSockets and native mobile apps."]
    ]
    add_table(doc, ["S.No", "Parameter", "Description"], psf)

    doc.save("docs/word_documents/08_Problem_Solution_Fit.docx")
    print("Generated: 08_Problem_Solution_Fit.docx")

# -------------------------------------------------------------
# 9. Proposed Solution
# -------------------------------------------------------------
def build_09():
    doc = create_base_doc("Proposed Solution", "Phase 4: Project Design Phase")
    add_meta_table(doc, "2 Marks")

    doc.add_heading("Executive Overview", level=2)
    doc.add_paragraph(
        "HotelPro solves the core vulnerabilities of legacy travel portals by unifying an overbooking-proof booking engine, "
        "transparent server pricing, automated tiered refunds, a front-desk staff portal, and a grounded AI concierge."
    )

    modules = [
        ["Module 1: Overbooking-Proof Engine", "Atomic 10-minute hold window (`Pending` state); real-time date overlap algorithm; automated 60s background cleanup worker."],
        ["Module 2: Pricing & Payments", "Authoritative server-side price computation; transparent 12% GST; Razorpay sandbox integration with HMAC signature verification."],
        ["Module 3: Grounded AI Concierge", "Google Gemini 1.5 Flash grounded in live catalog data; natural language search parsing; AI review pros/cons synthesis."],
        ["Module 4: Staff Operations Portal", "Lightweight front-desk portal; 1-click Check-in / Check-out; direct public manager review replies; room blackout dates."],
        ["Module 5: Admin Platform Governance", "Executive analytics dashboard visualizing revenue, bookings, and occupancy; role-based access control (`guest`, `staff`, `admin`)."]
    ]
    add_table(doc, ["Solution Module", "Engineering Capabilities & Deliverables"], modules)

    doc.save("docs/word_documents/09_Proposed_Solution.docx")
    print("Generated: 09_Proposed_Solution.docx")

# -------------------------------------------------------------
# 10. Solution Architecture
# -------------------------------------------------------------
def build_10():
    doc = create_base_doc("Solution Architecture", "Phase 4: Project Design Phase")
    add_meta_table(doc, "4 Marks")

    doc.add_heading("3-Tier Architecture Overview", level=2)
    doc.add_paragraph(
        "Tier 1 (Presentation): React 18 single page application with Vite, Tailwind CSS, and Axios interceptors.\n"
        "Tier 2 (Application Logic): Node.js + Express REST API with Zod validation, JWT dual-token auth, and BookingService.\n"
        "Tier 3 (Data & Services): MongoDB Atlas database, Razorpay test payment gateway, and Google Gemini AI API."
    )

    layers = [
        ["Presentation Layer (Client)", "React 18, Vite, Tailwind CSS, React Router v7", "Responsive UI, client routing, AuthContext, Axios 401 silent token renewal."],
        ["Middleware & Security", "Helmet, CORS, Cookie-Parser, Zod, JWT Guard", "Security headers, input sanitization, 15m access token + 7d httpOnly refresh cookie."],
        ["Application Controllers", "HotelCtrl, BookingCtrl, PaymentCtrl, StaffCtrl, AICtrl", "Decoupled HTTP request handling, standardized API response format."],
        ["Business Services", "BookingService, PaymentService, AIService", "Date overlap calculations, atomic holds, Razorpay HMAC verification, RAG prompts."],
        ["Persistence & External", "MongoDB Atlas, Razorpay Sandbox, Google Gemini", "Compound indexes on roomTypeId/status/dates; test checkout; catalog-grounded AI."]
    ]
    add_table(doc, ["Architectural Layer", "Technologies Selected", "Core Responsibilities"], layers)

    doc.save("docs/word_documents/10_Solution_Architecture.docx")
    print("Generated: 10_Solution_Architecture.docx")

# -------------------------------------------------------------
# 11. User Acceptance Testing (UAT)
# -------------------------------------------------------------
def build_11():
    doc = create_base_doc("User Acceptance Testing (UAT) Template", "Phase 5: Performance Testing Phase")
    add_meta_table(doc, "4 Marks")

    doc.add_heading("Testing Environment & Credentials", level=2)
    doc.add_paragraph(
        "Application URL: http://localhost:5173 | API Base: http://localhost:5001/api\n"
        "Admin: admin@hotelpro.com / Admin@123\n"
        "Staff: staff@hotelpro.com / Staff@123\n"
        "Guest: guest@hotelpro.com / Guest@123"
    )

    tc = [
        ["TC-001", "User Registration", "Submit name, email, password 'Secret123'", "Account created, strength bar valid, JWT issued", "Account created, logged in", "PASS"],
        ["TC-002", "Brute-Force Lockout", "Enter wrong password 5 consecutive times", "Account locked for 15 mins (HTTP 423)", "Locked on 5th attempt", "PASS"],
        ["TC-003", "Silent Token Renewal", "Let 15m access token expire; make request", "Axios interceptor renews token via cookie", "Renewed silently in background", "PASS"],
        ["TC-004", "Date Range Validation", "Select check-out before check-in or >30 nights", "Blocked with clear error message", "Validation error displayed", "PASS"],
        ["TC-005", "Live Availability Query", "Search hotels for specific date range", "Sold-out hotels excluded from results", "Only available hotels listed", "PASS"],
        ["TC-006", "10-Minute Hold Window", "Select room and click Reserve", "Hold created; 10m timer starts; vacant count drops", "10m timer displayed", "PASS"],
        ["TC-007", "Concurrency Collision Test", "20 concurrent requests for last room", "Exactly 1 booking succeeds; 19 rejected", "Exactly 1 hold created", "PASS"],
        ["TC-008", "Server Price Breakdown", "Select 2 nights @ 5500 + WELCOME10 coupon", "Subtotal 11k, -1.1k coupon, +12% GST = 11,088", "Matches server computation", "PASS"],
        ["TC-009", "Razorpay Test Payment", "Complete test card/UPI checkout", "HMAC verified; status Confirmed; voucher issued", "Signature verified, confirmed", "PASS"],
        ["TC-010", "Cancellation (>48h)", "Cancel booking with check-in 5 days away", "Status Cancelled; 100% full refund issued", "100% refund confirmed", "PASS"],
        ["TC-011", "Cancellation (24-48h)", "Cancel booking 30 hours before check-in", "Status Cancelled; 50% refund issued", "50% refund confirmed", "PASS"],
        ["TC-012", "Verified Review", "Submit review on completed stay", "Review saved; hotel average rating updates", "Review published, rating updated", "PASS"],
        ["TC-013", "Staff Check-in", "Click Check In on confirmed guest", "Status updates Confirmed -> Checked-in", "Updated to Checked-in", "PASS"],
        ["TC-014", "Admin Analytics", "Open /admin portal", "Revenue, bookings count, occupancy % render", "Dashboard KPIs displayed", "PASS"],
        ["TC-015", "AI Concierge Inquiry", "Ask: 'Which Goa hotel has pool and pets?'", "Accurately names Grand Azure Resort with rate", "Answered accurately from catalog", "PASS"]
    ]
    add_table(doc, ["Test ID", "Scenario", "Steps", "Expected Result", "Actual Result", "Status"], tc)

    doc.save("docs/word_documents/11_User_Acceptance_Testing_UAT.docx")
    print("Generated: 11_User_Acceptance_Testing_UAT.docx")

if __name__ == "__main__":
    build_01()
    build_02()
    build_03()
    build_04()
    build_05()
    build_06()
    build_07()
    build_08()
    build_09()
    build_10()
    build_11()
    print("\n✅ ALL 11 WORD DOCUMENTS GENERATED SUCCESSFULLY IN docs/word_documents/!")
