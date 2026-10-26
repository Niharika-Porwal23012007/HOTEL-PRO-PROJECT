import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def style_paragraph(p, text):
    # parse bold **text** and `code`
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Courier New'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(194, 65, 12)
        else:
            p.add_run(part)

def md_to_docx(md_path, docx_path, title_text=None):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    def flush_table(lines_to_flush):
        if not lines_to_flush:
            return
        parsed_rows = []
        for l in lines_to_flush:
            # strip start/end pipes
            stripped = l.strip()
            if stripped.startswith('|'):
                stripped = stripped[1:]
            if stripped.endswith('|'):
                stripped = stripped[:-1]
            cells = [c.strip() for c in stripped.split('|')]
            # check if separator
            if all(set(c.strip()).issubset({'-', ':', ' '}) for c in cells if c.strip()):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        cols = max(len(r) for r in parsed_rows)
        # normalize row lengths
        for r in parsed_rows:
            while len(r) < cols:
                r.append("")

        table = doc.add_table(rows=len(parsed_rows), cols=cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_data in enumerate(parsed_rows):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            bg = "1E3A8A" if is_header else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for c_idx, cell_value in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                style_paragraph(p, cell_value)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if is_header:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
        doc.add_paragraph()

    def flush_code(lines_to_flush):
        if not lines_to_flush:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.2)
        code_text = "".join(lines_to_flush)
        r = p.add_run(code_text)
        r.font.name = 'Courier New'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        # Background shading for code block
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        p._p.get_or_add_pPr().append(shd)

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        if not stripped:
            continue

        # Handle Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(stripped[2:])
            r.bold = True
            r.font.size = Pt(18)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped[3:])
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(30, 58, 138)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(stripped[4:])
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[5:])
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(71, 85, 105)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            style_paragraph(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            prefix = re.match(r'^\d+\.\s', stripped).group(0)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            style_paragraph(p, stripped[len(prefix):])
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped[2:])
            r.italic = True
            r.font.color.rgb = RGBColor(75, 85, 99)
        elif stripped == '---':
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            style_paragraph(p, stripped)

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Created docx: {docx_path}")

print("Converter defined.")
